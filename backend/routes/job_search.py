"""Feature 26 — Job Search. AI-powered job search workspace.

Replaces the legacy Jobs Portal (feature_id: jobs-portal, feature_number: 26).
Workflow: Career Profile -> Search & Match (AI fit score) -> Document Studio
(CV + cover letter, drafter-reviewer loop) -> ATS validation -> Tracker.
"""

import json
import logging
import os
import re
import uuid
from datetime import datetime, timedelta, timezone
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel

from routes.db import db, require_auth, require_admin, EMERGENT_LLM_KEY
from emergentintegrations.llm.chat import LlmChat, UserMessage

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/job-search")

VALID_STATUSES = ["saved", "applied", "interview", "offer", "rejected"]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _parse_llm_json(text: str):
    clean = (text or "").strip()
    if clean.startswith("```json"):
        clean = clean[7:]
    if clean.startswith("```"):
        clean = clean[3:]
    if clean.endswith("```"):
        clean = clean[:-3]
    clean = clean.strip()
    try:
        return json.loads(clean)
    except json.JSONDecodeError:
        match = re.search(r"\{[\s\S]*\}", clean)
        if match:
            return json.loads(match.group(0))
        raise


def _job_public(doc: dict) -> dict:
    return {
        "job_id": doc.get("job_id") or doc.get("slug") or "",
        "title": doc.get("title") or "",
        "department": doc.get("department") or "",
        "location": doc.get("location") or "",
        "type": doc.get("type") or "",
        "level": doc.get("level") or "",
        "description": doc.get("description") or "",
        "requirements": doc.get("requirements") or [],
        "salary_usd_min": doc.get("salary_usd_min"),
        "salary_usd_max": doc.get("salary_usd_max"),
        "posted_at": doc.get("posted_at") or "",
    }


async def _get_open_job(job_id: str) -> dict:
    doc = await db.careers_jobs.find_one({"job_id": job_id, "status": "open"})
    if not doc:
        doc = await db.careers_jobs.find_one({"slug": job_id, "status": "open"})
    if not doc:
        raise HTTPException(status_code=404, detail="Job not found or no longer open")
    return doc


# ── External job portals (repo-faithful registry + extensions) ──

JOB_PORTALS = [
    {"label": "Akademikernes Jobbank", "domains": ["jobbank.dk"]},
    {"label": "Jobdanmark.dk", "domains": ["jobdanmark.dk"]},
    {"label": "Jobindex.dk", "domains": ["jobindex.dk"]},
    {"label": "Jobnet.dk", "domains": ["jobnet.dk"]},
    {"label": "LinkedIn", "domains": ["linkedin.com"]},
    {"label": "freehire.dev", "domains": ["freehire.dev"]},
    {"label": "Indeed", "domains": ["indeed.com", "indeed.dk", "indeed.co.uk", "indeed.de", "indeed.fr"]},
    {"label": "Glassdoor", "domains": ["glassdoor.com", "glassdoor.dk", "glassdoor.co.uk", "glassdoor.de"]},
    {"label": "StepStone", "domains": ["stepstone.dk", "stepstone.de", "stepstone.com"]},
    {"label": "Monster", "domains": ["monster.com", "monster.dk", "monster.co.uk"]},
    {"label": "ZipRecruiter", "domains": ["ziprecruiter.com"]},
    {"label": "Welcome to the Jungle", "domains": ["welcometothejungle.com"]},
    {"label": "EURES", "domains": ["eures.europa.eu", "europa.eu"]},
]


def _identify_portal(url: str) -> str:
    try:
        from urllib.parse import urlparse

        host = (urlparse(url).hostname or "").lower().lstrip("www.")
        for portal in JOB_PORTALS:
            if any(host == d or host.endswith("." + d) for d in portal["domains"]):
                return portal["label"]
    except Exception:
        pass
    return "Generic"


def _html_to_text(html: str) -> str:
    import html as html_lib

    text = re.sub(r"(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>", " ", html)
    text = re.sub(r"(?i)<(br|/p|/div|/li|/h[1-6]|/tr)[^>]*>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html_lib.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


async def _extract_external_job(user_id: str, url: str, portal: str, raw_text: str) -> dict:
    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=f"jobimport_{uuid.uuid4().hex[:8]}",
        system_message=(
            "You are a job-posting parser. Extract the structured job posting from raw page text. "
            "If the text is clearly NOT a job posting, return {\"is_job_posting\": false}. Otherwise return ONLY a JSON object: "
            "{\"is_job_posting\": true, \"title\": string, \"company\": string, \"location\": string, "
            "\"employment_type\": string, \"level\": string, \"description\": string (concise 2-4 paragraph summary of the role), "
            "\"requirements\": [string], \"salary_text\": string, \"deadline\": string}"
        ),
    ).with_model("openai", "gpt-5.2")
    response = await chat.send_message(UserMessage(text=f"SOURCE URL: {url}\n\nPAGE TEXT:\n{raw_text[:15000]}"))
    parsed = _parse_llm_json(response.text if hasattr(response, "text") else str(response))
    if not parsed.get("is_job_posting"):
        raise HTTPException(status_code=422, detail="The page does not look like a job posting")

    doc = {
        "external_job_id": f"ext_{uuid.uuid4().hex[:12]}",
        "user_id": user_id,
        "source_url": url,
        "portal": portal,
        "title": str(parsed.get("title") or "")[:200],
        "company": str(parsed.get("company") or "")[:200],
        "location": str(parsed.get("location") or "")[:200],
        "employment_type": str(parsed.get("employment_type") or "")[:100],
        "level": str(parsed.get("level") or "")[:100],
        "description": str(parsed.get("description") or "")[:6000],
        "requirements": [str(r)[:300] for r in (parsed.get("requirements") or [])][:30],
        "salary_text": str(parsed.get("salary_text") or "")[:200],
        "deadline": str(parsed.get("deadline") or "")[:100],
        "imported_at": _now_iso(),
    }
    if not doc["title"]:
        raise HTTPException(status_code=422, detail="Could not extract a job title from the page")
    await db.job_search_external_jobs.insert_one(dict(doc))
    return doc


def _external_public(doc: dict) -> dict:
    return {
        "job_id": doc.get("external_job_id"),
        "title": doc.get("title") or "",
        "department": doc.get("company") or "",
        "location": doc.get("location") or "",
        "type": doc.get("employment_type") or "",
        "level": doc.get("level") or "",
        "description": doc.get("description") or "",
        "requirements": doc.get("requirements") or [],
        "salary_usd_min": None,
        "salary_usd_max": None,
        "salary_text": doc.get("salary_text") or "",
        "deadline": doc.get("deadline") or "",
        "posted_at": doc.get("imported_at") or "",
        "source_url": doc.get("source_url") or "",
        "portal": doc.get("portal") or "",
        "is_external": True,
    }


async def _resolve_job(user_id: str, job_id: str) -> dict:
    """Resolve internal (careers_jobs) or user-imported external (ext_*) jobs to a common shape."""
    if str(job_id).startswith("ext_"):
        doc = await db.job_search_external_jobs.find_one(
            {"external_job_id": job_id, "user_id": user_id}, {"_id": 0}
        )
        if not doc:
            raise HTTPException(status_code=404, detail="Imported job not found")
        public = _external_public(doc)
        public["job_id"] = job_id
        return public
    return _job_public(await _get_open_job(job_id))


def _job_context_block(job: dict) -> str:
    reqs = job.get("requirements") or []
    if isinstance(reqs, list):
        reqs = "; ".join(str(r) for r in reqs)
    salary = job.get("salary_text") or f"{job.get('salary_usd_min')}-{job.get('salary_usd_max')} USD"
    return (
        f"JOB POSTING\nTitle: {job.get('title')}\nCompany/Department: {job.get('department')}\n"
        f"Location: {job.get('location')}\nType: {job.get('type')}\nLevel: {job.get('level')}\n"
        f"Description: {job.get('description')}\nRequirements: {reqs}\n"
        f"Salary: {salary}"
    )


def _profile_context_block(profile: dict) -> str:
    return (
        f"CANDIDATE PROFILE\nHeadline: {profile.get('headline')}\nSummary: {profile.get('summary')}\n"
        f"Skills: {', '.join(profile.get('skills') or [])}\n"
        f"Experience: {profile.get('experience_years')} years\nEducation: {profile.get('education')}\n"
        f"Target roles: {', '.join(profile.get('target_roles') or [])}\n"
        f"Preferred location: {profile.get('preferred_location')} (remote preferred: {profile.get('remote_preferred')})\n"
        f"Key achievements: {profile.get('achievements')}"
    )


async def _require_profile(user_id: str) -> dict:
    profile = await db.job_search_profiles.find_one({"user_id": user_id}, {"_id": 0})
    if not profile or not (profile.get("summary") or profile.get("skills")):
        raise HTTPException(status_code=400, detail="Complete your career profile first")
    return profile


# ── Career Profile ────────────────────────────────────────


class CareerProfile(BaseModel):
    headline: str = ""
    summary: str = ""
    skills: List[str] = []
    experience_years: Optional[int] = None
    education: str = ""
    target_roles: List[str] = []
    preferred_location: str = ""
    remote_preferred: bool = False
    achievements: str = ""


@router.get("/profile")
async def get_profile(request: Request):
    user = await require_auth(request)
    profile = await db.job_search_profiles.find_one({"user_id": user.user_id}, {"_id": 0})
    return {"profile": profile}


@router.put("/profile")
async def save_profile(payload: CareerProfile, request: Request):
    user = await require_auth(request)
    doc = payload.model_dump()
    doc["skills"] = [s.strip() for s in doc.get("skills") or [] if str(s).strip()][:40]
    doc["target_roles"] = [s.strip() for s in doc.get("target_roles") or [] if str(s).strip()][:10]
    doc["user_id"] = user.user_id
    doc["updated_at"] = _now_iso()
    await db.job_search_profiles.update_one(
        {"user_id": user.user_id}, {"$set": doc, "$setOnInsert": {"created_at": _now_iso()}}, upsert=True
    )
    return {"success": True, "profile": doc}


# ── Job Search ────────────────────────────────────────────


@router.get("/jobs")
async def search_jobs(request: Request, q: str = "", location: str = "", job_type: str = "", limit: int = 50):
    await require_auth(request)
    query: dict = {"status": "open"}
    if q.strip():
        safe = re.escape(q.strip())
        query["$or"] = [
            {"title": {"$regex": safe, "$options": "i"}},
            {"department": {"$regex": safe, "$options": "i"}},
            {"description": {"$regex": safe, "$options": "i"}},
            {"level": {"$regex": safe, "$options": "i"}},
        ]
    if location.strip():
        query["location"] = {"$regex": re.escape(location.strip()), "$options": "i"}
    if job_type.strip():
        query["type"] = {"$regex": re.escape(job_type.strip()), "$options": "i"}
    cursor = db.careers_jobs.find(query).sort("posted_at", -1).limit(max(1, min(int(limit), 100)))
    jobs = [_job_public(doc) async for doc in cursor]
    return {"jobs": jobs, "total": len(jobs)}


# ── Search Alerts (saved queries surfaced by the weekly digest) ──


class SearchAlertRequest(BaseModel):
    q: str = ""
    location: str = ""
    job_type: str = ""
    label: str = ""


def _alert_internal_query(alert: dict, since_iso: str) -> dict:
    query: dict = {"status": "open", "posted_at": {"$gte": since_iso}}
    q = str(alert.get("q") or "").strip()
    if q:
        safe = re.escape(q)
        query["$or"] = [
            {"title": {"$regex": safe, "$options": "i"}},
            {"department": {"$regex": safe, "$options": "i"}},
            {"description": {"$regex": safe, "$options": "i"}},
            {"level": {"$regex": safe, "$options": "i"}},
        ]
    location = str(alert.get("location") or "").strip()
    if location:
        query["location"] = {"$regex": re.escape(location), "$options": "i"}
    job_type = str(alert.get("job_type") or "").strip()
    if job_type:
        query["type"] = {"$regex": re.escape(job_type), "$options": "i"}
    return query


async def collect_search_alert_matches(user_id: str, update_last_run: bool = False) -> list:
    """New internal + external roles matching each saved alert since its last run."""
    alerts = await db.job_search_alerts.find({"user_id": user_id}, {"_id": 0}).to_list(20)
    now = _now_iso()
    out = []
    for alert in alerts:
        since = str(alert.get("last_run_at") or alert.get("created_at") or now)
        q = str(alert.get("q") or "").strip()

        internal_cursor = db.careers_jobs.find(_alert_internal_query(alert, since)).sort("posted_at", -1).limit(10)
        internal = [_job_public(doc) async for doc in internal_cursor]

        ext_query: dict = {"user_id": user_id, "imported_at": {"$gte": since}}
        if q:
            safe = re.escape(q)
            ext_query["$or"] = [
                {"title": {"$regex": safe, "$options": "i"}},
                {"company": {"$regex": safe, "$options": "i"}},
                {"description": {"$regex": safe, "$options": "i"}},
            ]
        ext_cursor = db.job_search_external_jobs.find(ext_query, {"_id": 0}).sort("imported_at", -1).limit(10)
        external = [_external_public(doc) async for doc in ext_cursor]

        out.append(
            {
                "alert_id": alert.get("alert_id"),
                "label": str(alert.get("label") or q or "Saved search"),
                "query": {"q": q, "location": str(alert.get("location") or ""), "job_type": str(alert.get("job_type") or "")},
                "since": since,
                "notify_email": alert.get("notify_email", True) is not False,
                "new_internal_roles": internal,
                "new_external_roles": external,
                "new_match_count": len(internal) + len(external),
            }
        )
        if update_last_run:
            await db.job_search_alerts.update_one({"alert_id": alert.get("alert_id")}, {"$set": {"last_run_at": now}})
    return out


@router.post("/alerts")
async def create_search_alert(payload: SearchAlertRequest, request: Request):
    user = await require_auth(request)
    q = payload.q.strip()
    location = payload.location.strip()
    job_type = payload.job_type.strip()
    if not (q or location or job_type):
        raise HTTPException(status_code=400, detail="Provide at least one search criterion")

    existing = await db.job_search_alerts.find_one(
        {"user_id": user.user_id, "q": q, "location": location, "job_type": job_type}, {"_id": 0}
    )
    if existing:
        return {"alert": existing, "created": False}

    count = await db.job_search_alerts.count_documents({"user_id": user.user_id})
    if count >= 10:
        raise HTTPException(status_code=400, detail="Alert limit reached (10). Delete one to add another.")

    now = _now_iso()
    doc = {
        "alert_id": f"jsa_{uuid.uuid4().hex[:12]}",
        "user_id": user.user_id,
        "q": q,
        "location": location,
        "job_type": job_type,
        "label": payload.label.strip() or q or location or job_type,
        "notify_email": True,
        "created_at": now,
        "last_run_at": now,
    }
    await db.job_search_alerts.insert_one(dict(doc))
    return {"alert": doc, "created": True}


@router.get("/alerts")
async def list_search_alerts(request: Request):
    user = await require_auth(request)
    alerts = await collect_search_alert_matches(user.user_id)
    return {"alerts": alerts, "total": len(alerts)}


@router.delete("/alerts/{alert_id}")
async def delete_search_alert(alert_id: str, request: Request):
    user = await require_auth(request)
    result = await db.job_search_alerts.delete_one({"alert_id": alert_id, "user_id": user.user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Alert not found")
    return {"deleted": True}


class SearchAlertUpdate(BaseModel):
    notify_email: Optional[bool] = None


@router.patch("/alerts/{alert_id}")
async def update_search_alert(alert_id: str, payload: SearchAlertUpdate, request: Request):
    user = await require_auth(request)
    updates: dict = {}
    if payload.notify_email is not None:
        updates["notify_email"] = payload.notify_email
    if not updates:
        raise HTTPException(status_code=400, detail="No fields to update")
    result = await db.job_search_alerts.update_one(
        {"alert_id": alert_id, "user_id": user.user_id}, {"$set": updates}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Alert not found")
    fresh = await db.job_search_alerts.find_one({"alert_id": alert_id}, {"_id": 0})
    return {"alert": fresh}


async def run_search_alert_instant_scan(trigger: str = "scheduled") -> dict:
    """Email users instantly when new roles match their saved search alerts.

    Only advances the last_run_at cursor for alerts actually emailed, so the
    weekly digest keeps surfacing matches for email-muted alerts.
    """
    user_ids = await db.job_search_alerts.distinct("user_id")
    now = _now_iso()
    scanned = 0
    matched_users = 0
    sent = 0
    skipped_unsubscribed = 0

    for uid in user_ids:
        scanned += 1
        user = await db.users.find_one(
            {"user_id": uid}, {"_id": 0, "email": 1, "name": 1, "full_name": 1}
        )
        email = str((user or {}).get("email") or "").strip()
        if not email:
            continue

        pref = await db.job_alert_preferences.find_one(
            {"$or": [{"user_id": uid}, {"email": email.lower()}]}, {"_id": 0, "subscribed": 1}
        )
        if pref and pref.get("subscribed") is False:
            skipped_unsubscribed += 1
            continue

        results = await collect_search_alert_matches(uid, update_last_run=False)
        sections = []
        notified_alert_ids = []
        for r in results:
            if r["new_match_count"] <= 0 or not r.get("notify_email", True):
                continue
            roles = []
            for j in (r["new_internal_roles"] + r["new_external_roles"]):
                meta = " · ".join(
                    x for x in [
                        str(j.get("department") or j.get("company") or "").strip(),
                        str(j.get("location") or "").strip(),
                    ] if x
                )
                roles.append({"title": str(j.get("title") or "New role"), "meta": meta})
            sections.append({"label": r["label"], "roles": roles})
            notified_alert_ids.append(r["alert_id"])

        total_new = sum(len(s["roles"]) for s in sections)
        if total_new == 0:
            continue

        matched_users += 1
        name = (user or {}).get("full_name") or (user or {}).get("name") or "there"
        frontend_base = os.environ.get("FRONTEND_BASE_URL", "")
        from routes.job_alerts import _unsub_token

        unsub_url = f"{frontend_base}/api/jobs/alerts/unsubscribe?email={email}&token={_unsub_token(email)}"

        try:
            from utils.email_service import send_catalog_template

            await send_catalog_template(
                recipient_email=email,
                template_key="search_alert_instant",
                recipient_name=name,
                user_name=name,
                alert_sections=sections,
                total_new=total_new,
                unsub_url=unsub_url,
            )
            sent += 1
            await db.job_search_alerts.update_many(
                {"alert_id": {"$in": notified_alert_ids}}, {"$set": {"last_run_at": now, "last_notified_at": now}}
            )
        except Exception as e:
            logger.error(f"[search-alert-scan] send failed for {email}: {e}")

    summary = {
        "trigger": trigger,
        "scanned_users": scanned,
        "matched_users": matched_users,
        "emails_sent": sent,
        "skipped_unsubscribed": skipped_unsubscribed,
        "ran_at": now,
    }
    await db.search_alert_notification_log.insert_one(dict(summary))
    return summary


@router.post("/alerts/scan-now")
async def trigger_search_alert_scan(request: Request):
    await require_admin(request)
    return await run_search_alert_instant_scan(trigger="manual_admin")


# ── External Job Import (paste a link from public job sites) ──


class ExternalImportRequest(BaseModel):
    url: str


class ExternalManualImportRequest(BaseModel):
    url: str
    pasted_text: str


def _validate_job_url(url: str) -> str:
    from urllib.parse import urlparse

    url = (url or "").strip()
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise HTTPException(status_code=400, detail="Provide a valid http(s) job posting URL")
    host = parsed.hostname.lower()
    if host in ("localhost", "127.0.0.1", "0.0.0.0") or host.endswith(".local") or host.endswith(".internal"):
        raise HTTPException(status_code=400, detail="URL host not allowed")
    return url


@router.get("/external/portals")
async def list_supported_portals(request: Request):
    await require_auth(request)
    return {"portals": [p["label"] for p in JOB_PORTALS] + ["Generic (any job site)"]}


@router.post("/external/import")
async def import_external_job(payload: ExternalImportRequest, request: Request):
    import httpx

    user = await require_auth(request)
    url = _validate_job_url(payload.url)
    portal = _identify_portal(url)

    page_text = ""
    try:
        async with httpx.AsyncClient(
            timeout=15.0,
            follow_redirects=True,
            headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml",
                "Accept-Language": "en,da;q=0.8",
            },
        ) as client:
            resp = await client.get(url)
            if resp.status_code == 200 and "text/html" in (resp.headers.get("content-type") or "text/html"):
                page_text = _html_to_text(resp.text)
    except Exception as e:
        logger.info(f"External job fetch failed for {portal}: {e}")

    if len(page_text) < 300:
        return {
            "success": False,
            "fetch_blocked": True,
            "portal": portal,
            "detail": "This site blocks automated fetching — paste the job description instead.",
        }

    job_doc = await _extract_external_job(user.user_id, url, portal, page_text)
    return {"success": True, "fetch_blocked": False, "job": _external_public(job_doc)}


@router.post("/external/import-manual")
async def import_external_job_manual(payload: ExternalManualImportRequest, request: Request):
    user = await require_auth(request)
    url = _validate_job_url(payload.url)
    text = (payload.pasted_text or "").strip()
    if len(text) < 120:
        raise HTTPException(status_code=400, detail="Paste the full job description (at least a few sentences)")
    portal = _identify_portal(url)
    job_doc = await _extract_external_job(user.user_id, url, portal, text[:20000])
    return {"success": True, "job": _external_public(job_doc)}


@router.get("/external/jobs")
async def list_external_jobs(request: Request):
    user = await require_auth(request)
    cursor = db.job_search_external_jobs.find({"user_id": user.user_id}, {"_id": 0}).sort("imported_at", -1).limit(50)
    return {"jobs": [_external_public(d) async for d in cursor]}


@router.delete("/external/jobs/{external_job_id}")
async def delete_external_job(external_job_id: str, request: Request):
    user = await require_auth(request)
    result = await db.job_search_external_jobs.delete_one(
        {"external_job_id": external_job_id, "user_id": user.user_id}
    )
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Imported job not found")
    return {"success": True}


# ── AI Fit Match ──────────────────────────────────────────


class MatchRequest(BaseModel):
    job_id: str


@router.post("/match")
async def evaluate_match(payload: MatchRequest, request: Request):
    user = await require_auth(request)
    profile = await _require_profile(user.user_id)
    job = await _resolve_job(user.user_id, payload.job_id)

    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"jobmatch_{uuid.uuid4().hex[:8]}",
            system_message=(
                "You are a rigorous career-fit evaluator. Score a candidate against a job posting on 5 dimensions "
                "(0-100 each): skills, experience, culture, location, career_alignment. Be honest — never inflate. "
                "Return ONLY a JSON object: {\"overall_score\": number, \"recommendation\": \"strong_apply|apply|stretch|skip\", "
                "\"dimensions\": [{\"key\": string, \"label\": string, \"score\": number, \"note\": string}], "
                "\"strengths\": [string], \"gaps\": [string], \"summary\": string}"
            ),
        ).with_model("openai", "gpt-5.2")
        prompt = f"{_profile_context_block(profile)}\n\n{_job_context_block(job)}\n\nEvaluate the fit."
        response = await chat.send_message(UserMessage(text=prompt))
        result = _parse_llm_json(response.text if hasattr(response, "text") else str(response))
    except Exception as e:
        logger.warning(f"Job match evaluation failed: {e}")
        raise HTTPException(status_code=500, detail="Fit evaluation unavailable, please retry")

    match_doc = {
        "user_id": user.user_id,
        "job_id": job.get("job_id"),
        "job_title": job.get("title"),
        "overall_score": result.get("overall_score", 0),
        "recommendation": result.get("recommendation", "apply"),
        "dimensions": result.get("dimensions", []),
        "strengths": result.get("strengths", []),
        "gaps": result.get("gaps", []),
        "summary": result.get("summary", ""),
        "created_at": _now_iso(),
    }
    await db.job_search_matches.update_one(
        {"user_id": user.user_id, "job_id": job.get("job_id")}, {"$set": match_doc}, upsert=True
    )
    return {"match": match_doc}


@router.get("/matches")
async def list_matches(request: Request):
    user = await require_auth(request)
    cursor = db.job_search_matches.find({"user_id": user.user_id}, {"_id": 0}).sort("created_at", -1).limit(50)
    return {"matches": [doc async for doc in cursor]}


# ── Document Studio (CV + Cover Letter, drafter-reviewer) ─


class GenerateDocsRequest(BaseModel):
    job_id: str


@router.post("/documents/generate")
async def generate_documents(payload: GenerateDocsRequest, request: Request):
    user = await require_auth(request)
    profile = await _require_profile(user.user_id)
    job = await _resolve_job(user.user_id, payload.job_id)
    display_name = getattr(user, "name", None) or getattr(user, "full_name", None) or "Candidate"

    context = f"{_profile_context_block(profile)}\n\n{_job_context_block(job)}"
    try:
        drafter = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"jobdraft_{uuid.uuid4().hex[:8]}",
            system_message=(
                "You are an expert CV and cover-letter writer. Draft a tailored CV and cover letter in clean Markdown. "
                "STRICT honesty rule: use ONLY facts from the candidate profile — never fabricate skills, employers, or metrics. "
                f"Candidate display name: {display_name}. "
                "Return ONLY a JSON object: {\"cv_markdown\": string, \"cover_letter_markdown\": string}"
            ),
        ).with_model("openai", "gpt-5.2")
        draft_resp = await drafter.send_message(UserMessage(text=f"{context}\n\nDraft the tailored CV and cover letter."))
        draft = _parse_llm_json(draft_resp.text if hasattr(draft_resp, "text") else str(draft_resp))

        reviewer = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"jobreview_{uuid.uuid4().hex[:8]}",
            system_message=(
                "You are a senior hiring reviewer. Critique the drafted CV and cover letter against the job posting: "
                "missed keywords, weak framing, generic language. Then return improved final versions. "
                "Keep the honesty rule — no invented experience. "
                "Return ONLY a JSON object: {\"cv_markdown\": string, \"cover_letter_markdown\": string, \"reviewer_notes\": [string]}"
            ),
        ).with_model("openai", "gpt-5.2")
        review_prompt = (
            f"{context}\n\nDRAFT CV:\n{draft.get('cv_markdown', '')}\n\nDRAFT COVER LETTER:\n"
            f"{draft.get('cover_letter_markdown', '')}\n\nCritique and return the revised final versions."
        )
        review_resp = await reviewer.send_message(UserMessage(text=review_prompt))
        final = _parse_llm_json(review_resp.text if hasattr(review_resp, "text") else str(review_resp))
    except Exception as e:
        logger.warning(f"Document generation failed: {e}")
        raise HTTPException(status_code=500, detail="Document generation unavailable, please retry")

    doc = {
        "user_id": user.user_id,
        "job_id": job.get("job_id"),
        "job_title": job.get("title"),
        "cv_markdown": final.get("cv_markdown") or draft.get("cv_markdown", ""),
        "cover_letter_markdown": final.get("cover_letter_markdown") or draft.get("cover_letter_markdown", ""),
        "reviewer_notes": final.get("reviewer_notes", []),
        "ats_report": None,
        "generated_at": _now_iso(),
    }
    await db.job_search_documents.update_one(
        {"user_id": user.user_id, "job_id": job.get("job_id")}, {"$set": doc}, upsert=True
    )

    try:
        user_email = getattr(user, "email", None)
        if user_email:
            from utils.email_templates import build_job_search_kit_ready_email
            from utils.email_service import send_email

            tpl = build_job_search_kit_ready_email(user_name=display_name, job_title=job.get("title") or "")
            await send_email(
                recipient_email=user_email,
                subject=tpl.subject,
                content=tpl.html,
                recipient_name=display_name,
                content_text=tpl.text,
                template_key="job_search_kit_ready",
            )
    except Exception as e:
        logger.warning(f"Job search kit email failed (non-blocking): {e}")

    return {"document": doc}


@router.get("/documents")
async def get_documents(request: Request, job_id: str = ""):
    user = await require_auth(request)
    if job_id:
        doc = await db.job_search_documents.find_one({"user_id": user.user_id, "job_id": job_id}, {"_id": 0})
        return {"documents": [doc] if doc else []}
    cursor = db.job_search_documents.find({"user_id": user.user_id}, {"_id": 0}).sort("generated_at", -1).limit(20)
    return {"documents": [d async for d in cursor]}


def _markdown_to_pdf(markdown_text: str, title: str) -> bytes:
    """Render document markdown into a clean, unbranded PDF (headings, bullets, bold/italic)."""
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

    styles = getSampleStyleSheet()
    body = ParagraphStyle("JSBody", parent=styles["Normal"], fontSize=10, leading=15, spaceAfter=6, textColor=HexColor("#1E293B"))
    h1 = ParagraphStyle("JSH1", parent=styles["Heading1"], fontSize=17, leading=21, spaceBefore=4, spaceAfter=8, textColor=HexColor("#0F172A"))
    h2 = ParagraphStyle("JSH2", parent=styles["Heading2"], fontSize=13, leading=17, spaceBefore=10, spaceAfter=5, textColor=HexColor("#0F766E"))
    h3 = ParagraphStyle("JSH3", parent=styles["Heading3"], fontSize=11, leading=15, spaceBefore=8, spaceAfter=4, textColor=HexColor("#0F172A"))

    def _inline(text: str) -> str:
        out = escape(text)
        out = re.sub(r"\*\*((?:[^*]|\*(?!\*))+)\*\*", r"<b>\1</b>", out)
        out = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<i>\1</i>", out)
        return out

    flowables = []
    bullet_buffer: list = []

    def _flush_bullets():
        nonlocal bullet_buffer
        if bullet_buffer:
            flowables.append(
                ListFlowable(
                    [ListItem(Paragraph(item, body), leftIndent=14) for item in bullet_buffer],
                    bulletType="bullet", start="•", leftIndent=14,
                )
            )
            bullet_buffer = []

    for raw_line in (markdown_text or "").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if not stripped:
            _flush_bullets()
            flowables.append(Spacer(1, 4))
            continue
        if stripped.startswith("### "):
            _flush_bullets()
            flowables.append(Paragraph(_inline(stripped[4:]), h3))
        elif stripped.startswith("## "):
            _flush_bullets()
            flowables.append(Paragraph(_inline(stripped[3:]), h2))
        elif stripped.startswith("# "):
            _flush_bullets()
            flowables.append(Paragraph(_inline(stripped[2:]), h1))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            bullet_buffer.append(_inline(stripped[2:]))
        elif set(stripped) <= {"-", "_", "="} and len(stripped) >= 3:
            _flush_bullets()
            flowables.append(Spacer(1, 6))
        else:
            _flush_bullets()
            flowables.append(Paragraph(_inline(stripped), body))
    _flush_bullets()

    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer, pagesize=letter,
        leftMargin=0.85 * inch, rightMargin=0.85 * inch,
        topMargin=0.8 * inch, bottomMargin=0.8 * inch,
        title=title,
    )
    pdf.build(flowables)
    return buffer.getvalue()


def _add_docx_inline_runs(paragraph, text: str) -> None:
    """Render **bold** and *italic* markdown spans as real Word runs."""
    tokens = re.split(r"(\*\*(?:[^*]|\*(?!\*))+\*\*|(?<!\*)\*[^*\n]+\*(?!\*))", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def _markdown_to_docx(markdown_text: str, title: str) -> bytes:
    """Render document markdown into a clean, unbranded Word (.docx) document."""
    import io

    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.core_properties.title = title
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for raw_line in (markdown_text or "").splitlines():
        stripped = raw_line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            _add_docx_inline_runs(document.add_heading("", level=3), stripped[4:])
        elif stripped.startswith("## "):
            _add_docx_inline_runs(document.add_heading("", level=2), stripped[3:])
        elif stripped.startswith("# "):
            _add_docx_inline_runs(document.add_heading("", level=1), stripped[2:])
        elif stripped.startswith("- ") or stripped.startswith("* "):
            _add_docx_inline_runs(document.add_paragraph("", style="List Bullet"), stripped[2:])
        elif set(stripped) <= {"-", "_", "="} and len(stripped) >= 3:
            continue
        else:
            _add_docx_inline_runs(document.add_paragraph(""), stripped)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@router.get("/documents/{job_id}/docx")
async def download_document_docx(job_id: str, request: Request, view: str = "cv"):
    from fastapi.responses import Response

    user = await require_auth(request)
    if view not in ("cv", "cover"):
        raise HTTPException(status_code=400, detail="view must be 'cv' or 'cover'")
    doc = await db.job_search_documents.find_one({"user_id": user.user_id, "job_id": job_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="No documents generated for this job")
    content = doc.get("cv_markdown") if view == "cv" else doc.get("cover_letter_markdown")
    if not content:
        raise HTTPException(status_code=404, detail="Requested document is empty")

    label = "CV" if view == "cv" else "Cover Letter"
    docx_bytes = _markdown_to_docx(content, f"{label} — {doc.get('job_title') or job_id}")
    filename = f"{'cv' if view == 'cv' else 'cover-letter'}-{job_id}.docx"
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/documents/{job_id}/pdf")
async def download_document_pdf(job_id: str, request: Request, view: str = "cv"):
    from fastapi.responses import Response

    user = await require_auth(request)
    if view not in ("cv", "cover"):
        raise HTTPException(status_code=400, detail="view must be 'cv' or 'cover'")
    doc = await db.job_search_documents.find_one({"user_id": user.user_id, "job_id": job_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="No documents generated for this job")
    content = doc.get("cv_markdown") if view == "cv" else doc.get("cover_letter_markdown")
    if not content:
        raise HTTPException(status_code=404, detail="Requested document is empty")

    label = "CV" if view == "cv" else "Cover Letter"
    pdf_bytes = _markdown_to_pdf(content, f"{label} — {doc.get('job_title') or job_id}")
    filename = f"{'cv' if view == 'cv' else 'cover-letter'}-{job_id}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/documents/{job_id}/email")
async def email_document_kit(job_id: str, request: Request):
    """Email the CV + cover-letter PDFs to the signed-in user as attachments."""
    import base64

    user = await require_auth(request)
    user_email = (getattr(user, "email", None) or "").strip()
    if not user_email:
        raise HTTPException(status_code=400, detail="Your account has no email address")
    doc = await db.job_search_documents.find_one({"user_id": user.user_id, "job_id": job_id}, {"_id": 0})
    if not doc or not (doc.get("cv_markdown") or doc.get("cover_letter_markdown")):
        raise HTTPException(status_code=404, detail="No documents generated for this job")

    job_title = doc.get("job_title") or job_id
    attachments = []
    if doc.get("cv_markdown"):
        cv_pdf = _markdown_to_pdf(doc["cv_markdown"], f"CV — {job_title}")
        attachments.append({
            "filename": f"cv-{job_id}.pdf",
            "content": base64.b64encode(cv_pdf).decode("ascii"),
            "content_type": "application/pdf",
        })
        cv_docx = _markdown_to_docx(doc["cv_markdown"], f"CV — {job_title}")
        attachments.append({
            "filename": f"cv-{job_id}.docx",
            "content": base64.b64encode(cv_docx).decode("ascii"),
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        })
    if doc.get("cover_letter_markdown"):
        cover_pdf = _markdown_to_pdf(doc["cover_letter_markdown"], f"Cover Letter — {job_title}")
        attachments.append({
            "filename": f"cover-letter-{job_id}.pdf",
            "content": base64.b64encode(cover_pdf).decode("ascii"),
            "content_type": "application/pdf",
        })
        cover_docx = _markdown_to_docx(doc["cover_letter_markdown"], f"Cover Letter — {job_title}")
        attachments.append({
            "filename": f"cover-letter-{job_id}.docx",
            "content": base64.b64encode(cover_docx).decode("ascii"),
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        })

    from utils.email_templates import build_job_search_kit_email
    from utils.email_service import send_email

    display_name = getattr(user, "name", None) or getattr(user, "full_name", None) or "there"
    tpl = build_job_search_kit_email(user_name=display_name, job_title=job_title)
    result = await send_email(
        recipient_email=user_email,
        subject=tpl.subject,
        content=tpl.html,
        recipient_name=display_name,
        content_text=tpl.text,
        template_key="job_search_kit_email",
        attachments=attachments,
    )
    if not (result or {}).get("success"):
        logger.warning(f"Job search kit email failed: {(result or {}).get('error')}")
        raise HTTPException(status_code=502, detail="Could not send the email right now, please retry")
    return {"success": True, "sent_to": user_email, "attachments": [a["filename"] for a in attachments]}


# ── ATS Validation ────────────────────────────────────────


class AtsCheckRequest(BaseModel):
    job_id: str


@router.post("/ats-check")
async def ats_check(payload: AtsCheckRequest, request: Request):
    user = await require_auth(request)
    job = await _resolve_job(user.user_id, payload.job_id)
    doc = await db.job_search_documents.find_one({"user_id": user.user_id, "job_id": job.get("job_id")}, {"_id": 0})
    if not doc or not doc.get("cv_markdown"):
        raise HTTPException(status_code=400, detail="Generate documents for this job first")

    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"atscheck_{uuid.uuid4().hex[:8]}",
            system_message=(
                "You are an ATS (Applicant Tracking System) validation engine. Compare a CV against a job posting. "
                "Extract the posting's required/preferred keywords and classify coverage. Check parseability signals "
                "(contact info present, clear section headers, sane reading order). "
                "Return ONLY a JSON object: {\"ats_score\": number (0-100), \"covered_keywords\": [string], "
                "\"missing_keywords\": [string], \"parseability\": {\"score\": number, \"issues\": [string]}, "
                "\"recommendations\": [string]}"
            ),
        ).with_model("openai", "gpt-5.2")
        prompt = f"{_job_context_block(job)}\n\nCV TEXT:\n{doc.get('cv_markdown')}\n\nRun the ATS validation."
        response = await chat.send_message(UserMessage(text=prompt))
        report = _parse_llm_json(response.text if hasattr(response, "text") else str(response))
    except Exception as e:
        logger.warning(f"ATS check failed: {e}")
        raise HTTPException(status_code=500, detail="ATS validation unavailable, please retry")

    report["checked_at"] = _now_iso()
    await db.job_search_documents.update_one(
        {"user_id": user.user_id, "job_id": job.get("job_id")}, {"$set": {"ats_report": report}}
    )
    return {"ats_report": report}


# ── Application Tracker ───────────────────────────────────


class TrackRequest(BaseModel):
    job_id: str
    status: str = "saved"
    notes: str = ""


class TrackUpdateRequest(BaseModel):
    status: Optional[str] = None
    notes: Optional[str] = None


@router.get("/applications")
async def list_applications(request: Request):
    user = await require_auth(request)
    cursor = db.job_search_applications.find({"user_id": user.user_id}, {"_id": 0}).sort("updated_at", -1).limit(100)
    return {"applications": [a async for a in cursor]}


@router.post("/applications")
async def track_application(payload: TrackRequest, request: Request):
    user = await require_auth(request)
    if payload.status not in VALID_STATUSES:
        raise HTTPException(status_code=400, detail=f"Invalid status. Use one of: {', '.join(VALID_STATUSES)}")
    job = await _resolve_job(user.user_id, payload.job_id)
    existing = await db.job_search_applications.find_one({"user_id": user.user_id, "job_id": job.get("job_id")})
    if existing:
        raise HTTPException(status_code=409, detail="Job is already in your tracker")
    app_doc = {
        "application_id": f"jsa_{uuid.uuid4().hex[:12]}",
        "user_id": user.user_id,
        "job_id": job.get("job_id"),
        "job_title": job.get("title"),
        "department": job.get("department"),
        "location": job.get("location"),
        "source_url": job.get("source_url") or "",
        "portal": job.get("portal") or "",
        "status": payload.status,
        "notes": payload.notes[:2000],
        "history": [{"status": payload.status, "at": _now_iso()}],
        "created_at": _now_iso(),
        "updated_at": _now_iso(),
    }
    await db.job_search_applications.insert_one(dict(app_doc))
    return {"success": True, "application": app_doc}


class MarkAppliedRequest(BaseModel):
    job_id: str


@router.post("/applications/mark-applied")
async def mark_application_applied(payload: MarkAppliedRequest, request: Request):
    """Idempotent funnel loop-closer: ensure the job is tracked with status >= applied.

    Creates the application as 'applied' if untracked, advances 'saved' -> 'applied',
    and never downgrades an application already at applied/interview/offer/rejected.
    """
    user = await require_auth(request)
    job = await _resolve_job(user.user_id, payload.job_id)
    existing = await db.job_search_applications.find_one(
        {"user_id": user.user_id, "job_id": job.get("job_id")}, {"_id": 0}
    )
    if existing:
        if existing.get("status") != "saved":
            return {"success": True, "status": existing.get("status"), "changed": False}
        await db.job_search_applications.update_one(
            {"application_id": existing["application_id"]},
            {
                "$set": {"status": "applied", "updated_at": _now_iso()},
                "$push": {"history": {"status": "applied", "at": _now_iso()}},
            },
        )
        return {"success": True, "status": "applied", "changed": True}

    app_doc = {
        "application_id": f"jsa_{uuid.uuid4().hex[:12]}",
        "user_id": user.user_id,
        "job_id": job.get("job_id"),
        "job_title": job.get("title"),
        "department": job.get("department"),
        "location": job.get("location"),
        "source_url": job.get("source_url") or "",
        "portal": job.get("portal") or "",
        "status": "applied",
        "notes": "",
        "history": [{"status": "applied", "at": _now_iso()}],
        "created_at": _now_iso(),
        "updated_at": _now_iso(),
    }
    await db.job_search_applications.insert_one(dict(app_doc))
    return {"success": True, "status": "applied", "changed": True}


@router.patch("/applications/{application_id}")
async def update_application(application_id: str, payload: TrackUpdateRequest, request: Request):
    user = await require_auth(request)
    existing = await db.job_search_applications.find_one(
        {"user_id": user.user_id, "application_id": application_id}
    )
    if not existing:
        raise HTTPException(status_code=404, detail="Application not found")
    updates: dict = {"updated_at": _now_iso()}
    push = None
    if payload.status is not None:
        if payload.status not in VALID_STATUSES:
            raise HTTPException(status_code=400, detail=f"Invalid status. Use one of: {', '.join(VALID_STATUSES)}")
        updates["status"] = payload.status
        push = {"history": {"status": payload.status, "at": _now_iso()}}
    if payload.notes is not None:
        updates["notes"] = payload.notes[:2000]
    op: dict = {"$set": updates}
    if push:
        op["$push"] = push
    await db.job_search_applications.update_one({"application_id": application_id}, op)
    doc = await db.job_search_applications.find_one({"application_id": application_id}, {"_id": 0})
    return {"success": True, "application": doc}


@router.delete("/applications/{application_id}")
async def delete_application(application_id: str, request: Request):
    user = await require_auth(request)
    result = await db.job_search_applications.delete_one(
        {"user_id": user.user_id, "application_id": application_id}
    )
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Application not found")
    return {"success": True}


# ── Weekly Digest ─────────────────────────────────────────


class DigestPreferenceRequest(BaseModel):
    subscribed: bool


@router.get("/digest/preference")
async def get_digest_preference(request: Request):
    user = await require_auth(request)
    pref = await db.job_alert_preferences.find_one({"user_id": user.user_id}, {"_id": 0, "subscribed": 1})
    return {"subscribed": True if pref is None else bool(pref.get("subscribed", True))}


@router.post("/digest/preference")
async def set_digest_preference(payload: DigestPreferenceRequest, request: Request):
    user = await require_auth(request)
    email = (getattr(user, "email", None) or "").strip().lower()
    await db.job_alert_preferences.update_one(
        {"user_id": user.user_id},
        {"$set": {"user_id": user.user_id, "email": email, "subscribed": payload.subscribed, "updated_at": _now_iso()}},
        upsert=True,
    )
    return {"success": True, "subscribed": payload.subscribed}


@router.post("/digest/send-now")
async def send_digest_now(request: Request):
    await require_admin(request)
    from scheduler_jobs.digests import _send_job_search_weekly_digest

    result = await _send_job_search_weekly_digest(db, trigger="admin_manual", force=True)
    return result


@router.get("/digest/runs")
async def list_digest_runs(request: Request):
    await require_admin(request)
    cursor = db.job_search_digest_runs.find({}, {"_id": 0}).sort("generated_at", -1).limit(10)
    return {"runs": [r async for r in cursor]}


# ── Workspace Summary ─────────────────────────────────────


@router.get("/summary")
async def workspace_summary(request: Request):
    user = await require_auth(request)
    profile = await db.job_search_profiles.find_one({"user_id": user.user_id}, {"_id": 0, "summary": 1, "skills": 1})
    open_jobs = await db.careers_jobs.count_documents({"status": "open"})
    matches = await db.job_search_matches.count_documents({"user_id": user.user_id})
    documents = await db.job_search_documents.count_documents({"user_id": user.user_id})
    tracked = await db.job_search_applications.count_documents({"user_id": user.user_id})
    interviews = await db.job_search_applications.count_documents({"user_id": user.user_id, "status": "interview"})
    offers = await db.job_search_applications.count_documents({"user_id": user.user_id, "status": "offer"})
    return {
        "profile_complete": bool(profile and (profile.get("summary") or profile.get("skills"))),
        "open_jobs": open_jobs,
        "matches": matches,
        "documents": documents,
        "tracked": tracked,
        "interviews": interviews,
        "offers": offers,
        "last_sync": _now_iso(),
    }


# ── Job Hunt Pulse (Home Dashboard card) ─────────────────


@router.get("/pulse")
async def job_hunt_pulse(request: Request):
    """Compact job-hunt snapshot for the Home Dashboard pulse card (mirrors weekly digest signals)."""
    user = await require_auth(request)
    profile = await db.job_search_profiles.find_one(
        {"user_id": user.user_id}, {"_id": 0, "skills": 1, "target_roles": 1, "summary": 1}
    )
    profile_complete = bool(profile and (profile.get("summary") or profile.get("skills")))
    week_start = (datetime.now(timezone.utc) - timedelta(days=7)).isoformat()

    status_counts = {s: 0 for s in VALID_STATUSES}
    async for row in db.job_search_applications.find({"user_id": user.user_id}, {"_id": 0, "status": 1}):
        s = str(row.get("status") or "")
        if s in status_counts:
            status_counts[s] += 1
    tracked = sum(status_counts.values())

    new_matching = []
    if profile_complete:
        keywords = [
            str(k).lower()
            for k in (profile.get("skills") or []) + (profile.get("target_roles") or [])
            if str(k).strip()
        ]
        if keywords:
            new_jobs = await db.careers_jobs.find(
                {"status": "open", "posted_at": {"$gte": week_start}},
                {"_id": 0, "job_id": 1, "title": 1, "department": 1, "description": 1, "requirements": 1, "level": 1},
            ).to_list(200)
            scored = []
            for job in new_jobs:
                reqs = job.get("requirements") or []
                if isinstance(reqs, list):
                    reqs = " ".join(str(r) for r in reqs)
                terms = f"{job.get('title', '')} {job.get('department', '')} {job.get('description', '')} {reqs} {job.get('level', '')}".lower()
                hits = sum(1 for kw in keywords if kw and kw in terms)
                if hits > 0:
                    scored.append((hits, job))
            scored.sort(key=lambda item: item[0], reverse=True)
            new_matching = [{"job_id": j.get("job_id"), "title": j.get("title")} for _, j in scored]

    best = await db.job_search_matches.find(
        {"user_id": user.user_id, "created_at": {"$gte": week_start}},
        {"_id": 0, "overall_score": 1, "job_title": 1},
    ).sort("overall_score", -1).limit(1).to_list(1)

    kits_week = await db.job_search_documents.count_documents(
        {"user_id": user.user_id, "generated_at": {"$gte": week_start}}
    )
    apps_week = await db.job_search_applications.count_documents(
        {"user_id": user.user_id, "created_at": {"$gte": week_start}}
    )

    return {
        "profile_complete": profile_complete,
        "status_counts": status_counts,
        "tracked": tracked,
        "new_matching_count": len(new_matching),
        "new_matching_top": new_matching[:3],
        "best_fit_score": int(best[0].get("overall_score") or 0) if best else 0,
        "best_fit_job_title": str(best[0].get("job_title") or "") if best else "",
        "kits_this_week": kits_week,
        "apps_this_week": apps_week,
        "generated_at": _now_iso(),
    }

