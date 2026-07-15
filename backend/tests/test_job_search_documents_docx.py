"""Feature 26 Job Search — Document .docx (Word) export tests.

Verifies:
  1. GET /api/job-search/documents/{job_id}/docx?view=cv returns 200,
     application/vnd.openxmlformats-officedocument.wordprocessingml.document,
     Content-Disposition attachment with cv-<job>.docx filename.
  2. python-docx can open the body; document contains Heading paragraphs and
     bold runs; no raw '**' or '# ' markdown symbols in visible text.
  3. view=cover succeeds; unauth => 401/403; nonexistent job => 404; view=bogus => 400.
  4. Regression: /pdf endpoint still 200 for both views; /email kit endpoint still 200.
"""

import io
import os

import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
assert BASE_URL, "REACT_APP_BACKEND_URL must be set"

ADMIN_EMAIL = "admin@realaicoach.app"
ADMIN_PASSWORD = "NewAdminPass2026!"
KNOWN_JOB_ID = "job_principal-fullstack-engineer"
DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
WRITE_HEADERS = {"X-Requested-With": "XMLHttpRequest", "Content-Type": "application/json"}


def _login(session: requests.Session, email: str, password: str) -> bool:
    r = session.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": email, "password": password},
        headers=WRITE_HEADERS,
        timeout=30,
    )
    return r.status_code == 200


@pytest.fixture(scope="module")
def admin_session() -> requests.Session:
    s = requests.Session()
    if not _login(s, ADMIN_EMAIL, ADMIN_PASSWORD):
        pytest.skip("Admin login failed (possibly rate-limited)")
    return s


@pytest.fixture(scope="module")
def ensure_doc(admin_session: requests.Session) -> str:
    r = admin_session.get(
        f"{BASE_URL}/api/job-search/documents?job_id={KNOWN_JOB_ID}", timeout=20
    )
    if r.status_code == 200 and (r.json().get("documents") or []):
        return KNOWN_JOB_ID
    g = admin_session.post(
        f"{BASE_URL}/api/job-search/documents/generate",
        json={"job_id": KNOWN_JOB_ID},
        headers=WRITE_HEADERS,
        timeout=240,
    )
    if g.status_code != 200:
        pytest.skip(f"Could not generate document for {KNOWN_JOB_ID}: {g.status_code}")
    return KNOWN_JOB_ID


# ─────────────────────────── happy path ───────────────────────────

class TestJobSearchDocumentDocx:
    def test_cv_docx_success_headers_and_mime(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/docx?view=cv",
            timeout=60,
        )
        assert r.status_code == 200, r.text[:400]
        ct = r.headers.get("content-type", "")
        assert ct.startswith(DOCX_MIME), ct
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower(), cd
        assert f"cv-{ensure_doc}.docx" in cd, cd
        # DOCX is a zip — starts with PK
        assert r.content[:2] == b"PK", r.content[:16]

    def test_cv_docx_python_docx_readback_and_content(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/docx?view=cv",
            timeout=60,
        )
        assert r.status_code == 200
        try:
            from docx import Document  # type: ignore
        except Exception:
            pytest.skip("python-docx not installed")
        doc = Document(io.BytesIO(r.content))
        # Must have at least one Heading paragraph
        heading_styles = [p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert heading_styles, "No Heading paragraphs found in docx"
        # Must have at least one bold run
        bold_runs = [
            run for p in doc.paragraphs for run in p.runs if run.bold
        ]
        assert bold_runs, "No bold runs found in docx"
        # No raw markdown tokens in visible text
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "**" not in full_text, "raw '**' bold markers leaked into docx text"
        assert "# " not in full_text, "raw '# ' heading markers leaked into docx text"
        # Content sanity: substantial text
        assert len(full_text.strip()) > 100, f"docx text too short ({len(full_text)} chars)"

    def test_cover_docx_success(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/docx?view=cover",
            timeout=60,
        )
        assert r.status_code == 200, r.text[:400]
        assert r.headers.get("content-type", "").startswith(DOCX_MIME)
        cd = r.headers.get("content-disposition", "")
        assert f"cover-letter-{ensure_doc}.docx" in cd, cd
        assert r.content[:2] == b"PK"

    def test_unauthenticated_401(self):
        r = requests.get(
            f"{BASE_URL}/api/job-search/documents/{KNOWN_JOB_ID}/docx?view=cv",
            timeout=15,
        )
        assert r.status_code in (401, 403), f"Expected 401/403 got {r.status_code}"

    def test_nonexistent_job_404(self, admin_session):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/job_does_not_exist_xyz/docx?view=cv",
            timeout=20,
        )
        assert r.status_code == 404, r.text[:200]

    def test_invalid_view_400(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/docx?view=bogus",
            timeout=20,
        )
        assert r.status_code == 400, r.text[:200]


# ─────────────────────────── regression ───────────────────────────

class TestJobSearchRegression:
    def test_pdf_cv_still_200(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/pdf?view=cv",
            timeout=60,
        )
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")

    def test_pdf_cover_still_200(self, admin_session, ensure_doc):
        r = admin_session.get(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/pdf?view=cover",
            timeout=60,
        )
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")

    def test_email_kit_still_200(self, admin_session, ensure_doc):
        # Endpoint name aligns with existing suite. Accept 200 or 202.
        r = admin_session.post(
            f"{BASE_URL}/api/job-search/documents/{ensure_doc}/email",
            json={},
            headers=WRITE_HEADERS,
            timeout=60,
        )
        assert r.status_code in (200, 202), f"Unexpected email kit status {r.status_code}: {r.text[:400]}"
